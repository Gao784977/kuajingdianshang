"""Loader for "keyword" Excel sheets (keyword aggregate exports).

Each row is one keyword with its search volume, growth, click /
conversion share, competition level and related keywords. One keyword
per row — duplicates across multiple sheets are merged downstream,
not here.
"""

from __future__ import annotations

import os
from typing import Any, Dict, List, Optional

from .column_aliases import build_column_mapping
from .normalized_models import KeywordMetric, to_decimal, to_str, to_str_list
from .product_loader import LoaderResult, _get_cell, _is_blank_row


KEYWORD_FIELDS: List[str] = [
    "keyword", "search_volume", "search_growth", "click_share",
    "conversion_share", "competition", "related_keywords",
]


def _build_keyword_from_row(
    row: List[Any],
    col_to_index: Dict[str, int],
    source_file: str,
    source_sheet: str,
    source_row: int,
) -> KeywordMetric:
    def field(name: str) -> Any:
        return _get_cell(row, col_to_index, name)

    return KeywordMetric(
        keyword=to_str(field("keyword")) or "",
        search_volume=to_decimal(field("search_volume")),
        search_growth=to_decimal(field("search_growth")),
        click_share=to_decimal(field("click_share")),
        conversion_share=to_decimal(field("conversion_share")),
        competition=to_str(field("competition")) or "",
        related_keywords=to_str_list(field("related_keywords")),
        source_file=source_file,
        source_sheet=source_sheet,
        source_row=source_row,
        source_type="imported_excel",
    )


def load(file_path: str, sheet_name: Optional[str] = None) -> LoaderResult:
    result = LoaderResult()
    if not file_path or not os.path.exists(file_path):
        result.errors.append(f"File not found: {file_path}")
        return result
    try:
        from openpyxl import load_workbook  # type: ignore[import]
    except ImportError as exc:
        result.errors.append(f"openpyxl not available: {exc}")
        return result
    try:
        wb = load_workbook(filename=file_path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        result.errors.append(f"Failed to open workbook: {exc}")
        return result
    try:
        if not wb.sheetnames:
            result.errors.append("Workbook has no sheets.")
            return result
        chosen_sheet = sheet_name if sheet_name in wb.sheetnames else wb.sheetnames[0]
        if sheet_name and sheet_name not in wb.sheetnames:
            result.warnings.append(
                f"Sheet '{sheet_name}' not found; falling back to '{chosen_sheet}'."
            )
        ws = wb[chosen_sheet]
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            result.errors.append("Sheet has no header row.")
            return result
        headers = [str(h).strip() if h is not None else "" for h in header_row]
        mapping = build_column_mapping(headers, "keyword")
        result.column_mapping = mapping
        col_to_index: Dict[str, int] = {}
        for m in mapping.mappings:
            if m.standard_field not in col_to_index:
                col_to_index[m.standard_field] = m.original_position
        records: List[Any] = []
        row_index = 1
        for row in rows_iter:
            row_index += 1
            row_list = list(row) if isinstance(row, tuple) else (list(row) if row else [])
            if _is_blank_row(row_list):
                continue
            try:
                kw = _build_keyword_from_row(
                    row_list,
                    col_to_index,
                    source_file=file_path,
                    source_sheet=chosen_sheet,
                    source_row=row_index,
                )
            except Exception as exc:  # noqa: BLE001
                result.errors.append(f"Row {row_index} failed to parse: {exc}")
                continue
            if not kw.keyword:
                result.warnings.append(f"Row {row_index} has no keyword; skipping.")
                continue
            records.append(kw)
        result.records = records
        result.row_count = len(records)
        if mapping.required_missing:
            result.warnings.append(
                "Missing required fields: " + ", ".join(mapping.required_missing)
            )
        if mapping.unmapped_columns:
            result.warnings.append(
                f"{len(mapping.unmapped_columns)} unmapped column(s): "
                + ", ".join(mapping.unmapped_columns[:5])
            )
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass
    return result


__all__ = ["load", "LoaderResult", "KEYWORD_FIELDS", "_build_keyword_from_row"]
