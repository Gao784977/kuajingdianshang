"""Loader for "brands" Excel sheets (seller-sprite brands list).

Distinct from :mod:`brand_analytics_loader`: this file is the seller
sprite "Brands" export where each row is one brand with monthly /
annual sales and revenue. No click / conversion share — those only
exist in Amazon Brand Analytics.
"""

from __future__ import annotations

import os
from typing import Any, Dict, List, Optional

from .column_aliases import build_column_mapping
from .normalized_models import BrandMetric, to_decimal, to_str
from .product_loader import LoaderResult, _get_cell, _is_blank_row


BRAND_FIELDS: List[str] = [
    "brand", "monthly_sales", "monthly_revenue", "annual_sales",
    "annual_revenue", "average_price", "market_share",
]


def _build_brand_from_row(
    row: List[Any],
    col_to_index: Dict[str, int],
    source_file: str,
    source_sheet: str,
    source_row: int,
) -> BrandMetric:
    def field(name: str) -> Any:
        return _get_cell(row, col_to_index, name)

    return BrandMetric(
        brand=to_str(field("brand")) or "",
        monthly_sales=to_decimal(field("monthly_sales")),
        monthly_revenue=to_decimal(field("monthly_revenue")),
        annual_sales=to_decimal(field("annual_sales")),
        annual_revenue=to_decimal(field("annual_revenue")),
        average_price=to_decimal(field("average_price")),
        market_share=to_decimal(field("market_share")),
        source_file=source_file,
        source_sheet=source_sheet,
        source_row=source_row,
        source_type="imported_excel",
    )


def load(file_path: str, sheet_name: Optional[str] = None) -> LoaderResult:
    result = LoaderResult()
    if not file_path or not os.path.exists(file_path):
        result.errors.append(f"File not found: {file_path}")
        return result
    try:
        from openpyxl import load_workbook  # type: ignore[import]
    except ImportError as exc:
        result.errors.append(f"openpyxl not available: {exc}")
        return result
    try:
        wb = load_workbook(filename=file_path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        result.errors.append(f"Failed to open workbook: {exc}")
        return result
    try:
        if not wb.sheetnames:
            result.errors.append("Workbook has no sheets.")
            return result
        chosen_sheet = sheet_name if sheet_name in wb.sheetnames else wb.sheetnames[0]
        if sheet_name and sheet_name not in wb.sheetnames:
            result.warnings.append(
                f"Sheet '{sheet_name}' not found; falling back to '{chosen_sheet}'."
            )
        ws = wb[chosen_sheet]
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            result.errors.append("Sheet has no header row.")
            return result
        headers = [str(h).strip() if h is not None else "" for h in header_row]
        mapping = build_column_mapping(headers, "brands")
        result.column_mapping = mapping
        col_to_index: Dict[str, int] = {}
        for m in mapping.mappings:
            if m.standard_field not in col_to_index:
                col_to_index[m.standard_field] = m.original_position
        records: List[Any] = []
        row_index = 1
        for row in rows_iter:
            row_index += 1
            row_list = list(row) if isinstance(row, tuple) else (list(row) if row else [])
            if _is_blank_row(row_list):
                continue
            try:
                brand = _build_brand_from_row(
                    row_list,
                    col_to_index,
                    source_file=file_path,
                    source_sheet=chosen_sheet,
                    source_row=row_index,
                )
            except Exception as exc:  # noqa: BLE001
                result.errors.append(f"Row {row_index} failed to parse: {exc}")
                continue
            if not brand.brand:
                result.warnings.append(f"Row {row_index} has no brand; skipping.")
                continue
            records.append(brand)
        result.records = records
        result.row_count = len(records)
        if mapping.required_missing:
            result.warnings.append(
                "Missing required fields: " + ", ".join(mapping.required_missing)
            )
        if mapping.unmapped_columns:
            result.warnings.append(
                f"{len(mapping.unmapped_columns)} unmapped column(s): "
                + ", ".join(mapping.unmapped_columns[:5])
            )
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass
    return result


__all__ = ["load", "LoaderResult", "BRAND_FIELDS", "_build_brand_from_row"]
