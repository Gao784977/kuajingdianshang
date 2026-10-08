"""Loader for "products" Excel sheets (Seller Sprite / Brand Analytics).

Reads an ``.xlsx`` file, normalises headers via
:mod:`column_aliases`, and emits one :class:`AmazonProduct` per data
row. Each record carries ``source_file`` / ``source_sheet`` /
``source_row`` so downstream agents can distinguish imported rows
from mock / user input.

Openpyxl is imported lazily inside :func:`load` so importing this
module does not require the optional dependency at import time.
"""

from __future__ import annotations

import os
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any, Dict, List, Optional

from .column_aliases import ColumnMappingResult, build_column_mapping
from .normalized_models import AmazonProduct, to_bool, to_decimal, to_int, to_str, to_str_list


# Fields this loader knows how to populate on :class:`AmazonProduct`.
PRODUCT_FIELDS: List[str] = [
    "asin", "sku", "parent_asin", "brand", "brand_link", "title",
    "bullet_points", "product_url", "main_image_url", "category_path",
    "category", "main_category", "subcategory", "main_bsr",
    "subcategory_bsr", "monthly_sales", "monthly_sales_growth",
    "monthly_revenue", "child_sales", "child_revenue", "variation_count",
    "price", "prime_price", "coupon", "qa_count", "review_count",
    "monthly_new_reviews", "rating", "review_rate", "fba_fee",
    "margin_rate", "launch_date", "listing_age_days", "fulfillment_method",
    "sellers_count", "buybox_seller", "seller_location", "weight",
    "dimensions", "package_weight", "package_dimensions",
]


@dataclass
class LoaderResult:
    """Container returned by every loader."""

    records: List[Any] = field(default_factory=list)
    column_mapping: Optional[ColumnMappingResult] = None
    row_count: int = 0
    warnings: List[str] = field(default_factory=list)
    errors: List[str] = field(default_factory=list)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "row_count": self.row_count,
            "warnings": list(self.warnings),
            "errors": list(self.errors),
            "column_mapping": self.column_mapping.to_dict()
            if self.column_mapping
            else None,
            "records": [r.to_dict() if hasattr(r, "to_dict") else r for r in self.records],
        }


def _get_cell(
    row: List[Any],
    col_to_index: Dict[str, int],
    field_name: str,
) -> Any:
    """Return the cell value for ``field_name`` from ``row`` or ``None``."""
    idx = col_to_index.get(field_name)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _build_product_from_row(
    row: List[Any],
    col_to_index: Dict[str, int],
    source_file: str,
    source_sheet: str,
    source_row: int,
) -> AmazonProduct:
    """Construct an :class:`AmazonProduct` from a single Excel row."""
    def field(field_name: str) -> Any:
        return _get_cell(row, col_to_index, field_name)

    return AmazonProduct(
        asin=to_str(field("asin")) or "",
        sku=to_str(field("sku")) or "",
        parent_asin=to_str(field("parent_asin")) or "",
        brand=to_str(field("brand")) or "",
        brand_link=to_str(field("brand_link")) or "",
        title=to_str(field("title")) or "",
        bullet_points=to_str_list(field("bullet_points")),
        product_url=to_str(field("product_url")) or "",
        main_image_url=to_str(field("main_image_url")) or "",
        category_path=to_str_list(field("category_path")),
        category=to_str(field("category")) or "",
        main_category=to_str(field("main_category")) or "",
        subcategory=to_str(field("subcategory")) or "",
        main_bsr=to_int(field("main_bsr")),
        subcategory_bsr=to_int(field("subcategory_bsr")),
        monthly_sales=to_decimal(field("monthly_sales")),
        monthly_sales_growth=to_decimal(field("monthly_sales_growth")),
        monthly_revenue=to_decimal(field("monthly_revenue")),
        child_sales=to_decimal(field("child_sales")),
        child_revenue=to_decimal(field("child_revenue")),
        variation_count=to_int(field("variation_count")),
        price=to_decimal(field("price")),
        prime_price=to_decimal(field("prime_price")),
        coupon=to_decimal(field("coupon")),
        qa_count=to_int(field("qa_count")),
        review_count=to_int(field("review_count")),
        monthly_new_reviews=to_int(field("monthly_new_reviews")),
        rating=to_decimal(field("rating")),
        review_rate=to_decimal(field("review_rate")),
        fba_fee=to_decimal(field("fba_fee")),
        margin_rate=to_decimal(field("margin_rate")),
        launch_date=to_str(field("launch_date")),
        listing_age_days=to_int(field("listing_age_days")),
        fulfillment_method=to_str(field("fulfillment_method")) or "",
        sellers_count=to_int(field("sellers_count")),
        buybox_seller=to_str(field("buybox_seller")) or "",
        seller_location=to_str(field("seller_location")) or "",
        weight=to_decimal(field("weight")),
        dimensions=to_str(field("dimensions")) or "",
        package_weight=to_decimal(field("package_weight")),
        package_dimensions=to_str(field("package_dimensions")) or "",
        source_file=source_file,
        source_sheet=source_sheet,
        source_row=source_row,
        source_type="imported_excel",
    )


def _is_blank_row(row: List[Any]) -> bool:
    """True if every cell in ``row`` is None or empty string."""
    for v in row:
        if v is None:
            continue
        if isinstance(v, str) and not v.strip():
            continue
        return False
    return True


def load(file_path: str, sheet_name: Optional[str] = None) -> LoaderResult:
    """Load a products workbook and return typed records.

    Args:
        file_path: Path to the ``.xlsx`` file.
        sheet_name: Optional explicit sheet name. If ``None`` the
            first sheet is used.

    Returns:
        A :class:`LoaderResult`. The function never raises on a
        malformed cell — bad rows are skipped with a warning recorded
        in ``result.errors`` and ``result.warnings``.
    """
    result = LoaderResult()
    if not file_path or not os.path.exists(file_path):
        result.errors.append(f"File not found: {file_path}")
        return result
    try:
        from openpyxl import load_workbook  # type: ignore[import]
    except ImportError as exc:
        result.errors.append(f"openpyxl not available: {exc}")
        return result
    try:
        wb = load_workbook(filename=file_path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        result.errors.append(f"Failed to open workbook: {exc}")
        return result
    try:
        if not wb.sheetnames:
            result.errors.append("Workbook has no sheets.")
            return result
        chosen_sheet = sheet_name if sheet_name in wb.sheetnames else wb.sheetnames[0]
        if sheet_name and sheet_name not in wb.sheetnames:
            result.warnings.append(
                f"Sheet '{sheet_name}' not found; falling back to "
                f"'{chosen_sheet}'."
            )
        ws = wb[chosen_sheet]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter, None)
        except Exception as exc:  # noqa: BLE001
            result.errors.append(f"Failed to read header row: {exc}")
            return result
        if not header_row:
            result.errors.append("Sheet has no header row.")
            return result
        headers = [str(h).strip() if h is not None else "" for h in header_row]
        mapping = build_column_mapping(headers, "products")
        result.column_mapping = mapping
        # Map standard_field -> row index, using the first header that
        # matched that field.
        col_to_index: Dict[str, int] = {}
        for m in mapping.mappings:
            if m.standard_field not in col_to_index:
                col_to_index[m.standard_field] = m.original_position
        records: List[Any] = []
        row_index = 1  # 1-based; row 1 is the header.
        for row in rows_iter:
            row_index += 1
            row_list = list(row) if isinstance(row, tuple) else (list(row) if row else [])
            if _is_blank_row(row_list):
                continue
            try:
                product = _build_product_from_row(
                    row_list,
                    col_to_index,
                    source_file=file_path,
                    source_sheet=chosen_sheet,
                    source_row=row_index,
                )
            except Exception as exc:  # noqa: BLE001
                result.errors.append(
                    f"Row {row_index} failed to parse: {exc}"
                )
                continue
            # Skip rows that have no identifying info at all.
            if not (product.asin or product.title or product.brand):
                result.warnings.append(
                    f"Row {row_index} has no ASIN/title/brand; skipping."
                )
                continue
            records.append(product)
        result.records = records
        result.row_count = len(records)
        if mapping.required_missing:
            result.warnings.append(
                "Missing required fields: " + ", ".join(mapping.required_missing)
            )
        if mapping.unmapped_columns:
            result.warnings.append(
                f"{len(mapping.unmapped_columns)} unmapped column(s): "
                + ", ".join(mapping.unmapped_columns[:5])
            )
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass
    return result


__all__ = ["load", "LoaderResult", "PRODUCT_FIELDS"]
