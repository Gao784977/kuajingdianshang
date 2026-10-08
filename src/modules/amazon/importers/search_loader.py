"""Loader for "search" Excel sheets (search-term appearance / SERP scrape).

Each row captures one (keyword, ASIN) appearance pair alongside its
rank, page number, placement (sponsored vs organic) etc. Multiple
appearances of the same ASIN under the same keyword are intentionally
preserved — downstream agents can aggregate if needed, but the loader
must not dedup, otherwise the natural-vs-ad slot distinction is
lost.

Openpyxl is imported lazily inside :func:`load`.
"""

from __future__ import annotations

import os
from typing import Any, Dict, List, Optional

from .column_aliases import ColumnMappingResult, build_column_mapping
from .normalized_models import AmazonSearchAppearance, to_bool, to_decimal, to_int, to_str
from .product_loader import LoaderResult, _get_cell, _is_blank_row


SEARCH_FIELDS: List[str] = [
    "keyword", "asin", "search_rank", "placement_type", "placement_detail",
    "is_sponsored", "is_organic", "page_number", "brand", "title", "price",
]


def _build_search_from_row(
    row: List[Any],
    col_to_index: Dict[str, int],
    source_file: str,
    source_sheet: str,
    source_row: int,
) -> AmazonSearchAppearance:
    def field(name: str) -> Any:
        return _get_cell(row, col_to_index, name)

    is_sponsored = to_bool(field("is_sponsored"))
    is_organic = to_bool(field("is_organic"))
    placement_type_raw = to_str(field("placement_type")) or "unknown"
    # If the placement column is missing but the boolean flags exist,
    # infer placement_type from them so downstream consumers can group.
    if placement_type_raw == "unknown":
        if is_sponsored and not is_organic:
            placement_type_raw = "sponsored"
        elif is_organic and not is_sponsored:
            placement_type_raw = "organic"
    else:
        placement_type_raw = placement_type_raw.lower()
        if placement_type_raw not in ("sponsored", "organic", "unknown"):
            placement_type_raw = "unknown"

    return AmazonSearchAppearance(
        keyword=to_str(field("keyword")) or "",
        asin=to_str(field("asin")) or "",
        search_rank=to_int(field("search_rank")),
        placement_type=placement_type_raw,
        placement_detail=to_str(field("placement_detail")) or "",
        is_sponsored=is_sponsored,
        is_organic=is_organic,
        page_number=to_int(field("page_number")),
        brand=to_str(field("brand")) or "",
        title=to_str(field("title")) or "",
        price=to_decimal(field("price")),
        source_file=source_file,
        source_sheet=source_sheet,
        source_row=source_row,
        source_type="imported_excel",
    )


def load(file_path: str, sheet_name: Optional[str] = None) -> LoaderResult:
    """Load a search-appearance workbook.

    Rows are never deduped: the same ASIN may legitimately appear under
    multiple placements / pages / sponsored-vs-organic slots for the
    same keyword, and that multiplicity is exactly what downstream
    search-position analysis needs to consume.
    """
    result = LoaderResult()
    if not file_path or not os.path.exists(file_path):
        result.errors.append(f"File not found: {file_path}")
        return result
    try:
        from openpyxl import load_workbook  # type: ignore[import]
    except ImportError as exc:
        result.errors.append(f"openpyxl not available: {exc}")
        return result
    try:
        wb = load_workbook(filename=file_path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        result.errors.append(f"Failed to open workbook: {exc}")
        return result
    try:
        if not wb.sheetnames:
            result.errors.append("Workbook has no sheets.")
            return result
        chosen_sheet = sheet_name if sheet_name in wb.sheetnames else wb.sheetnames[0]
        if sheet_name and sheet_name not in wb.sheetnames:
            result.warnings.append(
                f"Sheet '{sheet_name}' not found; falling back to '{chosen_sheet}'."
            )
        ws = wb[chosen_sheet]
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            result.errors.append("Sheet has no header row.")
            return result
        headers = [str(h).strip() if h is not None else "" for h in header_row]
        mapping = build_column_mapping(headers, "search")
        result.column_mapping = mapping
        col_to_index: Dict[str, int] = {}
        for m in mapping.mappings:
            if m.standard_field not in col_to_index:
                col_to_index[m.standard_field] = m.original_position
        records: List[Any] = []
        row_index = 1
        for row in rows_iter:
            row_index += 1
            row_list = list(row) if isinstance(row, tuple) else (list(row) if row else [])
            if _is_blank_row(row_list):
                continue
            try:
                appearance = _build_search_from_row(
                    row_list,
                    col_to_index,
                    source_file=file_path,
                    source_sheet=chosen_sheet,
                    source_row=row_index,
                )
            except Exception as exc:  # noqa: BLE001
                result.errors.append(f"Row {row_index} failed to parse: {exc}")
                continue
            if not (appearance.keyword or appearance.asin):
                result.warnings.append(
                    f"Row {row_index} has no keyword/asin; skipping."
                )
                continue
            records.append(appearance)
        result.records = records
        result.row_count = len(records)
        if mapping.required_missing:
            result.warnings.append(
                "Missing required fields: " + ", ".join(mapping.required_missing)
            )
        if mapping.unmapped_columns:
            result.warnings.append(
                f"{len(mapping.unmapped_columns)} unmapped column(s): "
                + ", ".join(mapping.unmapped_columns[:5])
            )
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass
    return result


__all__ = ["load", "LoaderResult", "SEARCH_FIELDS", "_build_search_from_row"]
